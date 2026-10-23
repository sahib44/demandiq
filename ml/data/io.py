"""Strict file readers preserving source identifiers and duplicate headers."""

from __future__ import annotations

import csv
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class InputError(ValueError):
    """A supplied file is not a supported transaction table."""


def read_transactions(path: Path) -> pd.DataFrame:
    """Read XLSX or UTF-8 CSV without converting missing identifiers to strings.

    CSV fields remain strings for explicit normalization in the cleaner. The
    workbook's first sheet must have a single header row. Schema enforcement is
    shared with callers of the in-memory cleaner.
    """
    path = Path(path)
    try:
        if path.suffix.lower() == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = csv.reader(stream)
                headers = next(rows, None)
                if headers is None:
                    raise InputError("Input CSV has no header row")
                records = []
                for source_row, values in enumerate(rows, 1):
                    if len(values) != len(headers):
                        raise InputError(
                            f"CSV data row {source_row} has an unexpected field count"
                        )
                    records.append(values)
            return pd.DataFrame(records, columns=headers)
        if path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                rows = workbook.worksheets[0].iter_rows(values_only=True)
                headers = next(rows, None)
                if headers is None:
                    raise InputError("Input workbook has no header row")
                return pd.DataFrame(rows, columns=headers)
            finally:
                workbook.close()
        raise InputError("Supported input formats are .xlsx and UTF-8 .csv")
    except (
        OSError,
        UnicodeError,
        csv.Error,
        ValueError,
        zipfile.BadZipFile,
        ParseError,
        InvalidFileException,
    ) as exc:
        if isinstance(exc, InputError):
            raise
        # Do not include source values or full records in exception messages.
        raise InputError("Could not read transaction file; verify its format") from exc
